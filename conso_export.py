def consommation_export_excel(request):
    """Export Excel de la liste des consommations."""
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "CONSOMMATION_VIEW" not in permissions:
        messages.error(request, "Vous n'avez pas l'autorisation.")
        return redirect("consommations")

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    consommations_liste = (
        Consommation.objects
        .select_related(
            "id_detail_pec", "id_personne_beneficiaire", "id_adhesion",
            "id_acte", "id_sous_acte", "id_prestataire",
        )
        .order_by("-date_prestation", "-id_consommation")
    )

    workbook = openpyxl.Workbook()
    feuille = workbook.active
    feuille.title = "Consommations"

    font_titre = Font(bold=True, size=14, color="123B65")
    font_entete = Font(bold=True, color="FFFFFF", size=10)
    font_total = Font(bold=True, size=11, color="123B65")
    fill_entete = PatternFill(start_color="123B65", end_color="123B65", fill_type="solid")
    fill_total = PatternFill(start_color="EAF2FB", end_color="EAF2FB", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    feuille["A1"] = "LISTE DES CONSOMMATIONS"
    feuille["A1"].font = font_titre
    feuille["A1"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A1:L1")

    feuille["A2"] = f"Exporté le {timezone.now().strftime('%d/%m/%Y à %H:%M')}"
    feuille["A2"].font = Font(italic=True, size=9, color="666666")
    feuille["A2"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A2:L2")

    entetes = [
        "N°", "Bénéficiaire", "Acte", "Sous-acte", "Prestataire",
        "Date prestation", "Quantité", "Montant base", "PEC", "Reste",
        "Statut", "Exercice",
    ]

    ligne_entete = 4
    for col_num, entete in enumerate(entetes, start=1):
        cellule = feuille.cell(row=ligne_entete, column=col_num, value=entete)
        cellule.font = font_entete
        cellule.fill = fill_entete
        cellule.alignment = Alignment(horizontal="center", vertical="center")
        cellule.border = border

    ligne = ligne_entete + 1
    total_base = 0
    total_pec = 0
    total_reste = 0

    for idx, conso in enumerate(consommations_liste):
        ligne_courante = ligne + idx
        benef = conso.id_personne_beneficiaire
        sous_acte = conso.id_sous_acte.code_sous_acte if conso.id_sous_acte else "-"

        values = [
            conso.id_consommation,
            f"{benef.nom} {benef.prenom}" if benef else "-",
            f"{conso.id_acte.code_acte} - {conso.id_acte.libelle[:30]}" if conso.id_acte else "-",
            sous_acte,
            conso.id_prestataire.raison_sociale if conso.id_prestataire else "-",
            conso.date_prestation.strftime("%d/%m/%Y") if conso.date_prestation else "-",
            float(conso.quantite or 0),
            float(conso.montant_base or 0),
            float(conso.montant_prise_en_charge or 0),
            float(conso.montant_reste or 0),
            conso.statut,
            conso.exercice or "-",
        ]

        for col_num, val in enumerate(values, start=1):
            cellule = feuille.cell(row=ligne_courante, column=col_num, value=val)
            cellule.border = border
            cellule.alignment = Alignment(vertical="center", horizontal="center")
            if idx % 2 == 1:
                cellule.fill = fill_alt

        total_base += float(conso.montant_base or 0)
        total_pec += float(conso.montant_prise_en_charge or 0)
        total_reste += float(conso.montant_reste or 0)

    ligne_total = ligne + len(consommations_liste)
    feuille.cell(row=ligne_total, column=1, value="TOTAL").font = font_total
    feuille.merge_cells(start_row=ligne_total, start_column=1, end_row=ligne_total, end_column=7)
    feuille.cell(row=ligne_total, column=1).alignment = Alignment(horizontal="right", vertical="center")

    for col_num, val in enumerate([total_base, total_pec, total_reste], start=8):
        cellule = feuille.cell(row=ligne_total, column=col_num, value=val)
        cellule.font = font_total
        cellule.fill = fill_total
        cellule.border = border
        cellule.alignment = Alignment(horizontal="center", vertical="center")

    for col_num in [11, 12]:
        cellule = feuille.cell(row=ligne_total, column=col_num, value="")
        cellule.fill = fill_total
        cellule.border = border

    largeurs = [6, 22, 35, 15, 25, 14, 10, 14, 14, 14, 12, 10]
    for i, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[chr(64 + i)].width = largeur

    feuille.freeze_panes = "A5"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="consommations_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    )
    workbook.save(response)
    return response

