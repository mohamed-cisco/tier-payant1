def facture_export_excel(request):
    """Export Excel de la liste des factures."""
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "FACTURE_VIEW" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation d'exporter les factures."
        )
        return redirect("factures")

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    # Récupérer les factures
    factures_liste = (
        Facture.objects
        .select_related("id_prestataire")
        .order_by("-date_facture", "-id_facture")
    )

    # Créer le classeur
    workbook = openpyxl.Workbook()
    feuille = workbook.active
    feuille.title = "Factures"

    # Styles
    font_titre = Font(bold=True, size=14, color="123B65")
    font_entete = Font(bold=True, color="FFFFFF", size=11)
    font_total = Font(bold=True, size=11, color="123B65")
    fill_entete = PatternFill(start_color="123B65", end_color="123B65", fill_type="solid")
    fill_total = PatternFill(start_color="EAF2FB", end_color="EAF2FB", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    # Titre
    feuille["A1"] = "LISTE DES FACTURES"
    feuille["A1"].font = font_titre
    feuille["A1"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A1:I1")

    # Sous-titre : date d'export
    feuille["A2"] = f"Exporté le {timezone.now().strftime('%d/%m/%Y à %H:%M')}"
    feuille["A2"].font = Font(italic=True, size=9, color="666666")
    feuille["A2"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A2:I2")

    # En-têtes (ligne 4)
    entetes = [
        "N° Facture",
        "Prestataire",
        "Date facture",
        "Date réception",
        "Montant total",
        "Montant validé",
        "Montant rejeté",
        "Statut",
        "Date validation",
    ]

    ligne_entete = 4
    for col_num, entete in enumerate(entetes, start=1):
        cellule = feuille.cell(row=ligne_entete, column=col_num, value=entete)
        cellule.font = font_entete
        cellule.fill = fill_entete
        cellule.alignment = Alignment(horizontal="center", vertical="center")
        cellule.border = border

    # Données
    ligne = ligne_entete + 1
    total_montant = 0
    total_valide = 0
    total_rejete = 0

    for idx, facture in enumerate(factures_liste):
        ligne_courante = ligne + idx
        values = [
            facture.numero_facture,
            facture.id_prestataire.raison_sociale,
            facture.date_facture.strftime("%d/%m/%Y") if facture.date_facture else "-",
            facture.date_reception.strftime("%d/%m/%Y") if facture.date_reception else "-",
            float(facture.montant_total or 0),
            float(facture.montant_valide or 0),
            float(facture.montant_rejete or 0),
            facture.statut,
            facture.date_validation.strftime("%d/%m/%Y") if facture.date_validation else "-",
        ]

        for col_num, val in enumerate(values, start=1):
            cellule = feuille.cell(row=ligne_courante, column=col_num, value=val)
            cellule.border = border
            cellule.alignment = Alignment(vertical="center", horizontal="center")
            if idx % 2 == 1:
                cellule.fill = fill_alt

        total_montant += float(facture.montant_total or 0)
        total_valide += float(facture.montant_valide or 0)
        total_rejete += float(facture.montant_rejete or 0)

    # Ligne de totaux
    ligne_total = ligne + len(factures_liste)
    feuille.cell(row=ligne_total, column=1, value="TOTAL").font = font_total
    feuille.merge_cells(
        start_row=ligne_total, start_column=1,
        end_row=ligne_total, end_column=4
    )
    feuille.cell(row=ligne_total, column=1).alignment = Alignment(horizontal="right", vertical="center")

    for col_num, val in enumerate([total_montant, total_valide, total_rejete], start=5):
        cellule = feuille.cell(row=ligne_total, column=col_num, value=val)
        cellule.font = font_total
        cellule.fill = fill_total
        cellule.border = border
        cellule.alignment = Alignment(horizontal="center", vertical="center")

    for col_num in [8, 9]:
        cellule = feuille.cell(row=ligne_total, column=col_num, value="")
        cellule.fill = fill_total
        cellule.border = border

    # Ajuster la largeur des colonnes
    largeurs = [16, 25, 13, 14, 15, 15, 15, 14, 15]
    for i, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[chr(64 + i)].width = largeur

    # Figer l'en-tête
    feuille.freeze_panes = "A5"

    # Envoyer le fichier
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="factures_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    )
    workbook.save(response)
    return response

