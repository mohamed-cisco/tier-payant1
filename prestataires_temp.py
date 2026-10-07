def prestataire_import_excel(request):

    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list(
            "id_permission__code_permission",
            flat=True
        )
    )

    if "PRESTATAIRE_CREATE" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation d'importer les prestataires."
        )
        return redirect("prestataires")

    if request.method == "POST":

        fichier = request.FILES.get("fichier")

        if not fichier:

            messages.error(
                request,
                "Veuillez sélectionner un fichier Excel."
            )

            return redirect(
                "prestataire_import_excel"
            )

        if not fichier.name.endswith(".xlsx"):

            messages.error(
                request,
                "Le fichier doit être au format .xlsx"
            )

            return redirect(
                "prestataire_import_excel"
            )

        try:

            import openpyxl

            workbook = openpyxl.load_workbook(
                fichier
            )

            feuille = workbook.active

            nombre_importes = 0
            nombre_doublons = 0

            with transaction.atomic():

                for ligne in feuille.iter_rows(
                    min_row=2,
                    values_only=True
                ):

                    raison_sociale = ligne[1]

                    type_prestataire = ligne[2]

                    nif = ligne[3]

                    registre_commerce = ligne[4]

                    adresse = ligne[5]

                    telephone = ligne[6]

                    email = ligne[7]

                    statut = ligne[8]


                    if not raison_sociale:

                      continue


                    Prestataire.objects.create(

                        code_prestataire=_generer_code_prestataire(),

                        raison_sociale=str(
                            raison_sociale or ""
                        ).strip(),

                        type_prestataire=str(
                            type_prestataire or ""
                        ).strip(),

                        nif=str(
                            nif or ""
                        ).strip() or None,

                        registre_commerce=str(
                            registre_commerce or ""
                        ).strip() or None,

                        adresse=str(
                            adresse or ""
                        ).strip() or None,

                        telephone=str(
                            telephone or ""
                        ).strip() or None,

                        email=str(
                            email or ""
                        ).strip() or None,

                        statut=str(
                            statut or "ACTIF"
                        ).strip(),

                        date_creation=timezone.now(),

                    )

                    nombre_importes += 1


            messages.success(
                request,
                f"{nombre_importes} prestataire(s) importé(s) avec succès."
            )


            if nombre_doublons > 0:

                messages.warning(
                    request,
                    f"{nombre_doublons} doublon(s) ignoré(s)."
                )


            return redirect("prestataires")


        except Exception as e:

            messages.error(
                request,
                f"Erreur lors de l'importation : {str(e)}"
            )


    return render(
        request,
        "core/prestataire_import_excel.html",
        {
            "permissions": permissions,
            "page": "prestataires",
        }
    )
def prestataire_export_excel(request):

    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list(
            "id_permission__code_permission",
            flat=True
        )
    )

    if "PRESTATAIRE_VIEW" not in permissions:

        messages.error(
            request,
            "Vous n'avez pas l'autorisation d'exporter les prestataires."
        )

        return redirect("prestataires")


    import openpyxl

    from openpyxl.styles import Font

    from django.http import HttpResponse


    prestataires_list = (
        Prestataire.objects
        .all()
        .order_by("code_prestataire")
    )


    workbook = openpyxl.Workbook()

    feuille = workbook.active

    feuille.title = "Prestataires"


    entetes = [

        "Code prestataire",

        "Raison sociale",

        "Type prestataire",

        "NIF",

        "Registre de commerce",

        "Adresse",

        "Téléphone",

        "Email",

        "Statut",

        "Date création",

    ]


    feuille.append(entetes)


    for cellule in feuille[1]:

        cellule.font = Font(bold=True)


    for prestataire in prestataires_list:

        feuille.append([

            prestataire.code_prestataire,

            prestataire.raison_sociale,

            prestataire.type_prestataire,

            prestataire.nif,

            prestataire.registre_commerce,

            prestataire.adresse,

            prestataire.telephone,

            prestataire.email,

            prestataire.statut,

            prestataire.date_creation,

        ])


    for colonne in feuille.columns:

        longueur_max = 0

        lettre_colonne = colonne[0].column_letter


        for cellule in colonne:

            try:

                longueur = len(
                    str(cellule.value)
                ) if cellule.value else 0


                if longueur > longueur_max:

                    longueur_max = longueur

            except Exception:

                pass


        feuille.column_dimensions[
            lettre_colonne
        ].width = min(longueur_max + 2, 50)


    response = HttpResponse(

        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )

    )


    response[
        "Content-Disposition"
    ] = (
        'attachment; filename="prestataires.xlsx"'
    )


    workbook.save(response)


    return response

# ============================================================
# CHAMPS PERSONNALISÉS
# ============================================================

