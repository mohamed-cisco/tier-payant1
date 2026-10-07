def reglement_export_excel(request):
    """Export Excel de la liste des règlements."""
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "REGLEMENT_VIEW" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation d'exporter les règlements."
        )
        return redirect("reglements")

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    reglements_liste = (
        Reglement.objects
        .select_related("id_facture", "id_facture__id_prestataire")
        .order_by("-date_reglement", "-id_reglement")
    )

    workbook = openpyxl.Workbook()
    feuille = workbook.active
    feuille.title = "Reglements"

    font_titre = Font(bold=True, size=14, color="123B65")
    font_entete = Font(bold=True, color="FFFFFF", size=11)
    font_total = Font(bold=True, size=11, color="123B65")
    fill_entete = PatternFill(start_color="123B65", end_color="123B65", fill_type="solid")
    fill_total = PatternFill(start_color="EAF2FB", end_color="EAF2FB", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    # Titre
    feuille["A1"] = "LISTE DES RÈGLEMENTS"
    feuille["A1"].font = font_titre
    feuille["A1"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A1:I1")

    feuille["A2"] = f"Exporté le {timezone.now().strftime('%d/%m/%Y à %H:%M')}"
    feuille["A2"].font = Font(italic=True, size=9, color="666666")
    feuille["A2"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A2:I2")

    # En-têtes
    entetes = [
        "N° Règlement",
        "Référence",
        "N° Facture",
        "Prestataire",
        "Date règlement",
        "Montant",
        "Mode",
        "Statut",
        "Observation",
    ]

    ligne_entete = 4
    for col_num, entete in enumerate(entetes, start=1):
        cellule = feuille.cell(row=ligne_entete, column=col_num, value=entete)
        cellule.font = font_entete
        cellule.fill = fill_entete
        cellule.alignment = Alignment(horizontal="center", vertical="center")
        cellule.border = border

    # Données
    ligne = ligne_entete + 1
    total_montant = 0

    for idx, reglement in enumerate(reglements_liste):
        ligne_courante = ligne + idx
        values = [
            reglement.numero_reglement,
            reglement.reference_reglement or "-",
            reglement.id_facture.numero_facture,
            reglement.id_facture.id_prestataire.raison_sociale,
            reglement.date_reglement.strftime("%d/%m/%Y") if reglement.date_reglement else "-",
            float(reglement.montant or 0),
            reglement.mode_reglement or "-",
            reglement.statut,
            reglement.observation or "-",
        ]

        for col_num, val in enumerate(values, start=1):
            cellule = feuille.cell(row=ligne_courante, column=col_num, value=val)
            cellule.border = border
            cellule.alignment = Alignment(vertical="center", horizontal="center")
            if idx % 2 == 1:
                cellule.fill = fill_alt

        total_montant += float(reglement.montant or 0)

    # Total
    ligne_total = ligne + len(reglements_liste)
    feuille.cell(row=ligne_total, column=1, value="TOTAL").font = font_total
    feuille.merge_cells(
        start_row=ligne_total, start_column=1,
        end_row=ligne_total, end_column=5
    )
    feuille.cell(row=ligne_total, column=1).alignment = Alignment(horizontal="right", vertical="center")

    cellule = feuille.cell(row=ligne_total, column=6, value=total_montant)
    cellule.font = font_total
    cellule.fill = fill_total
    cellule.border = border
    cellule.alignment = Alignment(horizontal="center", vertical="center")

    for col_num in [7, 8, 9]:
        cellule = feuille.cell(row=ligne_total, column=col_num, value="")
        cellule.fill = fill_total
        cellule.border = border

    # Largeurs
    largeurs = [16, 18, 16, 25, 14, 15, 14, 14, 25]
    for i, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[chr(64 + i)].width = largeur

    feuille.freeze_panes = "A5"

    # Téléchargement
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="reglements_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    )
    workbook.save(response)
    return response

