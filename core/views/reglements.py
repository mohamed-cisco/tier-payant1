# core/views/reglements.py
"""
Vues de gestion des règlements.

Fonctions :
- reglements : liste
- reglement_create : créer
- reglement_detail : détail
- reglement_valider : valider
- reglement_export_excel : export Excel
- reglement_pdf : PDF
- _generer_numero_reglement : helper
- _generer_reference_reglement : helper
"""
from datetime import datetime, timedelta
from decimal import Decimal

import openpyxl
from openpyxl.styles import Font
from django.contrib import messages
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone

from core.forms import ReglementForm
from core.models import (
    Facture,
    Reglement,
    RolePermission,
)
from core.views.dashboard import enregistrer_audit

# Imports PDF
from django.conf import settings
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)


def _generer_reference_reglement(mode_reglement):
    """Génère une référence unique pour un règlement selon son mode."""
    annee = timezone.now().year

    prefixes = {
        "VIREMENT": f"VIR-{annee}-",
        "CHEQUE": f"CHQ-{annee}-",
        "ESPECES": f"ESP-{annee}-",
    }

    prefixe = prefixes.get(mode_reglement, f"REF-{annee}-")

    # Récupère toutes les références existantes pour ce préfixe
    references = (
        Reglement.objects
        .filter(reference_reglement__startswith=prefixe)
        .values_list("reference_reglement", flat=True)
    )

    valeurs = []
    for ref in references:
        try:
            valeurs.append(int(ref.rsplit("-", 1)[1]))
        except (ValueError, IndexError):
            continue

    prochain = max(valeurs, default=0) + 1
    reference = f"{prefixe}{prochain:04d}"

    # Vérifie l'unicité
    while Reglement.objects.filter(
        reference_reglement=reference
    ).exists():
        prochain += 1
        reference = f"{prefixe}{prochain:04d}"

    return reference
def _generer_numero_reglement():
    """Génère un numéro de règlement unique."""
    annee = timezone.now().year
    prefixe = f"REG-{annee}-"

    numeros = (
        Reglement.objects
        .filter(numero_reglement__startswith=prefixe)
        .values_list("numero_reglement", flat=True)
    )

    valeurs = []
    for numero in numeros:
        try:
            valeurs.append(int(numero.rsplit("-", 1)[1]))
        except (ValueError, IndexError):
            continue

    prochain = max(valeurs, default=0) + 1
    numero_reglement = f"{prefixe}{prochain:04d}"

    while Reglement.objects.filter(
        numero_reglement=numero_reglement
    ).exists():
        prochain += 1
        numero_reglement = f"{prefixe}{prochain:04d}"

    return numero_reglement






def reglement_create(request):
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "REGLEMENT_CREATE" not in permissions:
        messages.error(request, "Vous n'avez pas l'autorisation de créer un règlement.")
        return redirect("reglements")

    factures = (
        Facture.objects
        .filter(Q(statut="VALIDEE") | Q(statut="PARTIELLEMENT_PAYEE"))
        .select_related("id_prestataire")
        .order_by("-id_facture")
    )

    # Calculer le reste à payer pour chaque facture
    for f in factures:
        total_regle = (
            Reglement.objects
            .filter(id_facture=f, statut="VALIDEE")
            .aggregate(total=Sum("montant"))["total"]
            or 0
        )
        f.reste_a_payer = f.montant_valide - total_regle

    if request.method == "POST":
        form = ReglementForm(request.POST)
        print("=" * 50)
        print("🔍 [DEBUG] POST reçu")
        print("🔍 [DEBUG] Données POST :", dict(request.POST))
        print("🔍 [DEBUG] Form data :", form.data)

        form.fields["id_facture"].choices = [
            (
                str(f.id_facture),
                f"{f.numero_facture} - {f.id_prestataire.raison_sociale} - "
                f"Reste à payer : {f.reste_a_payer} DA"
            )
            for f in factures
        ]

        if form.is_valid():
            print("🔍 [DEBUG] Formulaire VALIDE")
            print("🔍 [DEBUG] cleaned_data :", form.cleaned_data)
            try:
                facture = Facture.objects.get(
                    id_facture=form.cleaned_data["id_facture"]
                )

                if facture.statut not in ["VALIDEE", "PARTIELLEMENT_PAYEE"]:
                    messages.error(request, "Cette facture ne peut plus être réglée.")
                    return redirect("reglements")

                montant = form.cleaned_data["montant"]

                montant_deja_regle = (
                    Reglement.objects
                    .filter(id_facture=facture, statut="VALIDEE")
                    .aggregate(total=Sum("montant"))["total"]
                    or 0
                )

                reste_a_payer = facture.montant_valide - montant_deja_regle

                if reste_a_payer <= 0:
                    messages.error(request, "Cette facture est déjà entièrement réglée.")
                    return render(
                        request,
                        "core/reglement_form.html",
                        {
                            "form": form,
                            "titre": "Nouveau règlement",
                            "factures": factures,
                        }
                    )

                if montant > reste_a_payer:
                    messages.error(
                        request,
                        f"Le montant du règlement ne peut pas dépasser "
                        f"le reste à payer de {reste_a_payer} DA."
                    )
                    return render(
                        request,
                        "core/reglement_form.html",
                        {
                            "form": form,
                            "titre": "Nouveau règlement",
                            "factures": factures,
                        }
                    )

                mode = form.cleaned_data["mode_reglement"]

                Reglement.objects.create(
                    id_facture=facture,
                    numero_reglement=_generer_numero_reglement(),
                    date_reglement=form.cleaned_data["date_reglement"],
                    montant=montant,
                    mode_reglement=mode,
                    reference_reglement=_generer_reference_reglement(mode),
                    statut="EN_ATTENTE",
                    observation=form.cleaned_data["observation"] or None,
                )

                messages.success(request, "Règlement créé avec succès.")
                return redirect("reglements")

            except Exception as e:
                messages.error(request, f"Erreur lors de la création du règlement : {e}")
                print("🔍 [DEBUG] ERREUR :", e)
    else:
        form = ReglementForm()

        form.fields["id_facture"].choices = [
            (
                str(f.id_facture),
                f"{f.numero_facture} - {f.id_prestataire.raison_sociale} - "
                f"Reste à payer : {f.reste_a_payer} DA"
            )
            for f in factures
        ]

    return render(
        request,
        "core/reglement_form.html",
        {
            "form": form,
            "titre": "Nouveau règlement",
            "factures": factures,
        }
    )





def reglements(request):
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list(
            "id_permission__code_permission",
            flat=True
        )
    )

    if "REGLEMENT_VIEW" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation de consulter les règlements."
        )
        return redirect("accueil")

    reglements = (
        Reglement.objects
        .select_related(
            "id_facture",
            "id_facture__id_prestataire",
        )
        .order_by("-id_reglement")
    )

    for reglement in reglements:
        montant_deja_regle = (
            Reglement.objects
            .filter(
                id_facture=reglement.id_facture,
                statut="VALIDEE"
            )
            .aggregate(total=Sum("montant"))["total"]
            or Decimal("0")
        )

        reglement.montant_total_facture = reglement.id_facture.montant_valide
        reglement.montant_deja_regle = montant_deja_regle
        reglement.reste_a_payer = (
            reglement.id_facture.montant_valide - montant_deja_regle
        )

    return render(
        request,
        "core/reglements.html",
        {
            "reglements": reglements,
            "permissions": permissions,
            "page": "reglements",     # ← AJOUTE
        }
    )





def reglement_export_excel(request):
    """Export Excel de la liste des règlements."""
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "REGLEMENT_VIEW" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation d'exporter les règlements."
        )
        return redirect("reglements")

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    reglements_liste = (
        Reglement.objects
        .select_related("id_facture", "id_facture__id_prestataire")
        .order_by("-date_reglement", "-id_reglement")
    )

    workbook = openpyxl.Workbook()
    feuille = workbook.active
    feuille.title = "Reglements"

    font_titre = Font(bold=True, size=14, color="123B65")
    font_entete = Font(bold=True, color="FFFFFF", size=11)
    font_total = Font(bold=True, size=11, color="123B65")
    fill_entete = PatternFill(start_color="123B65", end_color="123B65", fill_type="solid")
    fill_total = PatternFill(start_color="EAF2FB", end_color="EAF2FB", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    # Titre
    feuille["A1"] = "LISTE DES RÈGLEMENTS"
    feuille["A1"].font = font_titre
    feuille["A1"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A1:I1")

    feuille["A2"] = f"Exporté le {timezone.now().strftime('%d/%m/%Y à %H:%M')}"
    feuille["A2"].font = Font(italic=True, size=9, color="666666")
    feuille["A2"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A2:I2")

    # En-têtes
    entetes = [
        "N° Règlement",
        "Référence",
        "N° Facture",
        "Prestataire",
        "Date règlement",
        "Montant",
        "Mode",
        "Statut",
        "Observation",
    ]

    ligne_entete = 4
    for col_num, entete in enumerate(entetes, start=1):
        cellule = feuille.cell(row=ligne_entete, column=col_num, value=entete)
        cellule.font = font_entete
        cellule.fill = fill_entete
        cellule.alignment = Alignment(horizontal="center", vertical="center")
        cellule.border = border

    # Données
    ligne = ligne_entete + 1
    total_montant = 0

    for idx, reglement in enumerate(reglements_liste):
        ligne_courante = ligne + idx
        values = [
            reglement.numero_reglement,
            reglement.reference_reglement or "-",
            reglement.id_facture.numero_facture,
            reglement.id_facture.id_prestataire.raison_sociale,
            reglement.date_reglement.strftime("%d/%m/%Y") if reglement.date_reglement else "-",
            float(reglement.montant or 0),
            reglement.mode_reglement or "-",
            reglement.statut,
            reglement.observation or "-",
        ]

        for col_num, val in enumerate(values, start=1):
            cellule = feuille.cell(row=ligne_courante, column=col_num, value=val)
            cellule.border = border
            cellule.alignment = Alignment(vertical="center", horizontal="center")
            if idx % 2 == 1:
                cellule.fill = fill_alt

        total_montant += float(reglement.montant or 0)

    # Total
    ligne_total = ligne + len(reglements_liste)
    feuille.cell(row=ligne_total, column=1, value="TOTAL").font = font_total
    feuille.merge_cells(
        start_row=ligne_total, start_column=1,
        end_row=ligne_total, end_column=5
    )
    feuille.cell(row=ligne_total, column=1).alignment = Alignment(horizontal="right", vertical="center")

    cellule = feuille.cell(row=ligne_total, column=6, value=total_montant)
    cellule.font = font_total
    cellule.fill = fill_total
    cellule.border = border
    cellule.alignment = Alignment(horizontal="center", vertical="center")

    for col_num in [7, 8, 9]:
        cellule = feuille.cell(row=ligne_total, column=col_num, value="")
        cellule.fill = fill_total
        cellule.border = border

    # Largeurs
    largeurs = [16, 18, 16, 25, 14, 15, 14, 14, 25]
    for i, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[chr(64 + i)].width = largeur

    feuille.freeze_panes = "A5"

    # Téléchargement
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="reglements_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    )
    workbook.save(response)
    return response




def reglement_pdf(request, id_reglement):
    """Génère le PDF détaillé d'un règlement."""
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "REGLEMENT_VIEW" not in permissions:
        messages.error(request, "Vous n'avez pas l'autorisation.")
        return redirect("reglements")

    try:
        reglement = (
            Reglement.objects
            .select_related(
                "id_facture",
                "id_facture__id_prestataire",
            )
            .get(id_reglement=id_reglement)
        )
    except Reglement.DoesNotExist:
        messages.error(request, "Règlement introuvable.")
        return redirect("reglements")

    import os
    from django.conf import settings
    from django.http import HttpResponse
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image,
    )

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="Reglement-{reglement.numero_reglement}.pdf"'
    )

    document = SimpleDocTemplate(
        response, pagesize=A4,
        rightMargin=1.5 * cm, leftMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )

    elements = []
    styles = getSampleStyleSheet()

    style_titre = ParagraphStyle(
        "Titre", parent=styles["Title"], fontSize=20,
        textColor=colors.HexColor("#123b65"), alignment=2, leading=26,
    )
    style_section = ParagraphStyle(
        "Section", parent=styles["Heading2"], fontSize=12,
        textColor=colors.HexColor("#123b65"), spaceAfter=8, spaceBefore=10,
    )

    # En-tête
    logo_path = os.path.join(
        settings.BASE_DIR, "core", "static", "core", "img", "logo-sagps.png"
    )
    logo = Image(logo_path, width=5 * cm, height=2.2 * cm) if os.path.exists(logo_path) else ""

    titre_header = Paragraph(
        "<b>RÉCÉPISSÉ DE RÈGLEMENT</b><br/>"
        f"<font size=11>N° {reglement.numero_reglement}</font>",
        style_titre,
    )

    header_table = Table([[logo, titre_header]], colWidths=[7 * cm, 10.5 * cm])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LINEBELOW", (0, 0), (-1, 0), 2, colors.HexColor("#123b65")),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 0.5 * cm))

    # Infos règlement
    elements.append(Paragraph("INFORMATIONS DU RÈGLEMENT", style_section))

    infos_data = [
        ["N° Règlement", reglement.numero_reglement],
        ["Date", reglement.date_reglement.strftime("%d/%m/%Y") if reglement.date_reglement else "-"],
        ["Mode de règlement", reglement.mode_reglement or "-"],
        ["Référence", reglement.reference_reglement or "-"],
        ["Statut", reglement.statut],
    ]
    infos_table = Table(infos_data, colWidths=[4.5 * cm, 13 * cm])
    infos_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eaf2fb")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(infos_table)
    elements.append(Spacer(1, 0.5 * cm))

    # Prestataire
    elements.append(Paragraph("PRESTATAIRE", style_section))
    prest = reglement.id_facture.id_prestataire
    prest_data = [
        ["Raison sociale", prest.raison_sociale],
        ["Code prestataire", prest.code_prestataire],
        ["NIF", prest.nif or "-"],
        ["Adresse", prest.adresse or "-"],
    ]
    prest_table = Table(prest_data, colWidths=[4.5 * cm, 13 * cm])
    prest_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eaf2fb")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(prest_table)
    elements.append(Spacer(1, 0.5 * cm))

    # Facture
    elements.append(Paragraph("FACTURE CONCERNÉE", style_section))
    facture = reglement.id_facture
    facture_data = [
        ["N° Facture", facture.numero_facture],
        ["Date facture",
         facture.date_facture.strftime("%d/%m/%Y") if facture.date_facture else "-"],
        ["Montant total", f"{facture.montant_total:.2f} DA"],
        ["Montant validé", f"{facture.montant_valide:.2f} DA"],
    ]
    facture_table = Table(facture_data, colWidths=[4.5 * cm, 13 * cm])
    facture_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eaf2fb")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(facture_table)
    elements.append(Spacer(1, 0.5 * cm))

    # Montant payé (mis en évidence)
    elements.append(Paragraph("MONTANT RÉGLÉ", style_section))

    montant_data = [[
        "Montant",
        f"{reglement.montant:.2f} DA",
    ]]
    montant_table = Table(montant_data, colWidths=[6 * cm, 6 * cm], hAlign="CENTER")
    montant_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#123b65")),
        ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#dcfce7")),
        ("TEXTCOLOR", (0, 0), (0, 0), colors.white),
        ("TEXTCOLOR", (1, 0), (1, 0), colors.HexColor("#166534")),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 14),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
    ]))
    elements.append(montant_table)
    elements.append(Spacer(1, 0.8 * cm))

    # Observation
    if reglement.observation:
        elements.append(Paragraph("OBSERVATION", style_section))
        obs_table = Table([[reglement.observation]], colWidths=[17.5 * cm])
        obs_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        elements.append(obs_table)
        elements.append(Spacer(1, 0.5 * cm))

    # Signatures
    elements.append(Spacer(1, 1 * cm))
    sig_data = [[
        "Cachet du prestataire", "Cachet de l'organisme"
    ], [
        "\n\n\n_________________", "\n\n\n_________________",
    ]]
    sig_table = Table(sig_data, colWidths=[8.75 * cm, 8.75 * cm])
    sig_table.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, 0), 5),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 10),
    ]))
    elements.append(sig_table)

    document.build(elements)
    return response




def reglement_valider(request, id_reglement):
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list(
            "id_permission__code_permission",
            flat=True
        )
    )
    

    if "REGLEMENT_VALIDATE" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation de valider un règlement."
        )
        return redirect("reglements")

    try:
        reglement = Reglement.objects.get(
            id_reglement=id_reglement
        )
    except Reglement.DoesNotExist:
        messages.error(
            request,
            "Règlement introuvable."
        )
        return redirect("reglements")

    if request.method == "POST":

        if reglement.statut != "EN_ATTENTE":
            messages.error(
                request,
                "Ce règlement a déjÃ  été traité."
            )
            return redirect("reglements")

        reglement.statut = "VALIDEE"
        reglement.save()

        facture = reglement.id_facture

        montant_total_regle = (
              Reglement.objects
              .filter(
                  id_facture=facture,
                  statut="VALIDEE"
              )
              .aggregate(total=Sum("montant"))["total"]
              or 0
          )

        # Mettre à jour le statut de la facture selon le montant réglé
        if montant_total_regle >= facture.montant_valide:
            facture.statut = "PAYEE"
        elif montant_total_regle > 0:
            facture.statut = "PARTIELLEMENT_PAYEE"

        facture.save()

        messages.success(
            request,
            f"Règlement validé avec succès. "
            f"Facture {facture.numero_facture} : {facture.statut}."
        )

    return redirect("reglements")





def reglement_detail(request, id_reglement):
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list(
            "id_permission__code_permission",
            flat=True
        )
    )

    if "REGLEMENT_VIEW" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation de consulter le détail du règlement."
        )
        return redirect("reglements")

    try:
        reglement = (
            Reglement.objects
            .select_related(
                "id_facture",
                "id_facture__id_prestataire",
            )
            .get(id_reglement=id_reglement)
        )
    except Reglement.DoesNotExist:
        messages.error(
            request,
            "Règlement introuvable."
        )
        return redirect("reglements")

    return render(
        request,
        "core/reglement_detail.html",
        {
            "reglement": reglement,
            "page": "reglements",
        }
    )

