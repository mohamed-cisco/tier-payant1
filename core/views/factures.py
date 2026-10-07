# core/views/factures.py
"""
Vues de gestion des factures.

Fonctions :
- factures : liste
- facture_create : créer
- facture_detail : détail
- facture_valider : valider
- facture_export_excel : export Excel
- facture_export_pdf : export PDF
- facture_pdf : PDF facture
- _generer_numero_facture : helper
"""

from datetime import timedelta

from django.contrib import messages
from django.db.models import Q
from django.shortcuts import redirect, render
from django.utils import timezone

from core.forms import FactureForm
from core.models import (
    Consommation,
    DetailFacture,
    Facture,
    Prestataire,
    RolePermission,
)
from core.views.champs import (
    get_champs_pour_entite,
    get_valeurs_champs_entite,
    sauvegarder_valeurs_champs,
)
from core.views.dashboard import enregistrer_audit


def _generer_numero_facture():
    """Génère un numéro unique de facture."""
    annee = timezone.now().year
    prefixe = f"FAC-{annee}-"

    numeros = (
        Facture.objects
        .filter(numero_facture__startswith=prefixe)
        .values_list("numero_facture", flat=True)
    )

    valeurs = []
    for numero in numeros:
        try:
            valeurs.append(int(numero.rsplit("-", 1)[1]))
        except (ValueError, IndexError):
            continue

    prochain = max(valeurs, default=0) + 1
    numero_facture = f"{prefixe}{prochain:04d}"

    while Facture.objects.filter(numero_facture=numero_facture).exists():
        prochain += 1
        numero_facture = f"{prefixe}{prochain:04d}"

    return numero_facture


def factures(request):
    """Liste des factures."""
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "FACTURE_VIEW" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation de consulter les factures."
        )
        return redirect("accueil")

    factures_liste = (
        Facture.objects
        .select_related("id_prestataire")
        .order_by("-id_facture")
    )

    return render(
        request,
        "core/factures.html",
        {
            "factures": factures_liste,
            "permissions": permissions,
            "page": "factures",
        }
    )


def facture_create(request):
    """Créer une nouvelle facture."""
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "FACTURE_CREATE" not in permissions:
        messages.error(request, "Vous n'avez pas l'autorisation.")
        return redirect("factures")

    prestataires = (
        Prestataire.objects
        .filter(statut="ACTIF")
        .order_by("raison_sociale")
    )

    consommations = (
        Consommation.objects
        .filter(statut="VALIDEE")
        .select_related("id_prestataire", "id_acte", "id_personne_beneficiaire")
        .order_by("date_prestation", "id_consommation")
    )

    if request.method == "POST":
        form = FactureForm(request.POST)

        form.fields["id_prestataire"].choices = [
            (str(p.id_prestataire), f"{p.code_prestataire} - {p.raison_sociale}")
            for p in prestataires
        ]

        if form.is_valid():
            try:
                id_prestataire = form.cleaned_data["id_prestataire"]
                prestataire = Prestataire.objects.get(
                    id_prestataire=id_prestataire, statut="ACTIF"
                )

                # Détection de doublon
                date_limite = timezone.now().date() - timedelta(days=30)

                doublon = Facture.objects.filter(
                    id_prestataire=prestataire,
                    date_facture__gte=date_limite,
                    statut__in=["EN_ATTENTE", "VALIDEE"],
                ).order_by("-date_facture").first()

                if doublon and not request.POST.get("confirmer_doublon"):
                    messages.warning(
                        request,
                        f"⚠️ Une facture récente existe déjà pour ce prestataire : "
                        f"{doublon.numero_facture} "
                        f"({doublon.date_facture.strftime('%d/%m/%Y')}, "
                        f"{doublon.montant_valide} DA, statut {doublon.statut}). "
                        f"Voulez-vous vraiment créer une nouvelle facture ?"
                    )
                    return render(
                        request,
                        "core/facture_form.html",
                        {
                            "form": form,
                            "titre": "Nouvelle facture",
                            "consommations": consommations,
                            "page": "factures",
                            "doublon_detecte": doublon,
                            "champs_disponibles": get_champs_pour_entite("FACTURE"),
                        }
                    )

                # Traiter les consommations
                consommations_prestataire = [
                    c for c in consommations if c.id_prestataire_id == int(id_prestataire)
                ]

                consommations_deja_facturees = set(
                    DetailFacture.objects.values_list("id_consommation_id", flat=True)
                )

                consommations_prestataire = [
                    c for c in consommations_prestataire
                    if c.id_consommation not in consommations_deja_facturees
                    and c.montant_prise_en_charge > 0
                ]

                if not consommations_prestataire:
                    messages.error(
                        request,
                        "Aucune consommation disponible pour ce prestataire."
                    )
                    return render(
                        request,
                        "core/facture_form.html",
                        {
                            "form": form,
                            "titre": "Nouvelle facture",
                            "consommations": consommations,
                            "page": "factures",
                            "champs_disponibles": get_champs_pour_entite("FACTURE"),
                        }
                    )

                montant_total = sum(c.montant_base for c in consommations_prestataire)
                montant_valide = sum(c.montant_prise_en_charge for c in consommations_prestataire)
                montant_rejete = sum(c.montant_reste for c in consommations_prestataire)

                facture = Facture.objects.create(
                    id_prestataire=prestataire,
                    numero_facture=_generer_numero_facture(),
                    date_facture=form.cleaned_data["date_facture"],
                    date_reception=form.cleaned_data["date_reception"],
                    montant_total=montant_total,
                    montant_valide=montant_valide,
                    montant_rejete=montant_rejete,
                    statut="EN_ATTENTE",
                    date_validation=None,
                    utilisateur_validation=None,
                    observation=form.cleaned_data["observation"] or None,
                )

                for consommation in consommations_prestataire:
                    DetailFacture.objects.create(
                        id_facture=facture,
                        id_consommation=consommation,
                        id_acte=consommation.id_acte,
                        quantite=consommation.quantite,
                        montant_unitaire=(consommation.montant_base / consommation.quantite),
                        montant_total=consommation.montant_base,
                        montant_valide=consommation.montant_prise_en_charge,
                        montant_rejete=consommation.montant_reste,
                        statut="EN_ATTENTE",
                        id_motif_rejet=None,
                    )

                sauvegarder_valeurs_champs(request, "FACTURE", facture.id_facture)

                enregistrer_audit(
                    request=request,
                    type_action="CREATION",
                    module="FACTURE",
                    table_cible="facture",
                    id_enregistrement=facture.id_facture,
                    nouvelle_valeur=facture.numero_facture,
                    description=f"Création de la facture {facture.numero_facture}",
                )

                messages.success(
                    request,
                    f"Facture {facture.numero_facture} créée avec succès."
                )
                return redirect("factures")

            except Exception as e:
                messages.error(request, f"Erreur : {e}")
    else:
        form = FactureForm()
        form.fields["id_prestataire"].choices = [
            (str(p.id_prestataire), f"{p.code_prestataire} - {p.raison_sociale}")
            for p in prestataires
        ]

    champs = get_champs_pour_entite("FACTURE")
    for c in champs:
        c.valeur_actuelle = None
        c.choix_possibles_list = [
            x.strip() for x in (c.choix_possibles or "").split("\n") if x.strip()
        ]

    return render(
        request,
        "core/facture_form.html",
        {
            "form": form,
            "titre": "Nouvelle facture",
            "consommations": consommations,
            "page": "factures",
            "champs_disponibles": champs,
        }
    )


def facture_export_excel(request):
    """Export Excel de la liste des factures."""
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "FACTURE_VIEW" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation d'exporter les factures."
        )
        return redirect("factures")

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    # Récupérer les factures
    factures_liste = (
        Facture.objects
        .select_related("id_prestataire")
        .order_by("-date_facture", "-id_facture")
    )

    # Créer le classeur
    workbook = openpyxl.Workbook()
    feuille = workbook.active
    feuille.title = "Factures"

    # Styles
    font_titre = Font(bold=True, size=14, color="123B65")
    font_entete = Font(bold=True, color="FFFFFF", size=11)
    font_total = Font(bold=True, size=11, color="123B65")
    fill_entete = PatternFill(start_color="123B65", end_color="123B65", fill_type="solid")
    fill_total = PatternFill(start_color="EAF2FB", end_color="EAF2FB", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    # Titre
    feuille["A1"] = "LISTE DES FACTURES"
    feuille["A1"].font = font_titre
    feuille["A1"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A1:I1")

    # Sous-titre : date d'export
    feuille["A2"] = f"Exporté le {timezone.now().strftime('%d/%m/%Y à %H:%M')}"
    feuille["A2"].font = Font(italic=True, size=9, color="666666")
    feuille["A2"].alignment = Alignment(horizontal="center")
    feuille.merge_cells("A2:I2")

    # En-têtes (ligne 4)
    entetes = [
        "N° Facture",
        "Prestataire",
        "Date facture",
        "Date réception",
        "Montant total",
        "Montant validé",
        "Montant rejeté",
        "Statut",
        "Date validation",
    ]

    ligne_entete = 4
    for col_num, entete in enumerate(entetes, start=1):
        cellule = feuille.cell(row=ligne_entete, column=col_num, value=entete)
        cellule.font = font_entete
        cellule.fill = fill_entete
        cellule.alignment = Alignment(horizontal="center", vertical="center")
        cellule.border = border

    # Données
    ligne = ligne_entete + 1
    total_montant = 0
    total_valide = 0
    total_rejete = 0

    for idx, facture in enumerate(factures_liste):
        ligne_courante = ligne + idx
        values = [
            facture.numero_facture,
            facture.id_prestataire.raison_sociale,
            facture.date_facture.strftime("%d/%m/%Y") if facture.date_facture else "-",
            facture.date_reception.strftime("%d/%m/%Y") if facture.date_reception else "-",
            float(facture.montant_total or 0),
            float(facture.montant_valide or 0),
            float(facture.montant_rejete or 0),
            facture.statut,
            facture.date_validation.strftime("%d/%m/%Y") if facture.date_validation else "-",
        ]

        for col_num, val in enumerate(values, start=1):
            cellule = feuille.cell(row=ligne_courante, column=col_num, value=val)
            cellule.border = border
            cellule.alignment = Alignment(vertical="center", horizontal="center")
            if idx % 2 == 1:
                cellule.fill = fill_alt

        total_montant += float(facture.montant_total or 0)
        total_valide += float(facture.montant_valide or 0)
        total_rejete += float(facture.montant_rejete or 0)

    # Ligne de totaux
    ligne_total = ligne + len(factures_liste)
    feuille.cell(row=ligne_total, column=1, value="TOTAL").font = font_total
    feuille.merge_cells(
        start_row=ligne_total, start_column=1,
        end_row=ligne_total, end_column=4
    )
    feuille.cell(row=ligne_total, column=1).alignment = Alignment(horizontal="right", vertical="center")

    for col_num, val in enumerate([total_montant, total_valide, total_rejete], start=5):
        cellule = feuille.cell(row=ligne_total, column=col_num, value=val)
        cellule.font = font_total
        cellule.fill = fill_total
        cellule.border = border
        cellule.alignment = Alignment(horizontal="center", vertical="center")

    for col_num in [8, 9]:
        cellule = feuille.cell(row=ligne_total, column=col_num, value="")
        cellule.fill = fill_total
        cellule.border = border

    # Ajuster la largeur des colonnes
    largeurs = [16, 25, 13, 14, 15, 15, 15, 14, 15]
    for i, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[chr(64 + i)].width = largeur

    # Figer l'en-tête
    feuille.freeze_panes = "A5"

    # Envoyer le fichier
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="factures_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    )
    workbook.save(response)
    return response




def facture_valider(request, id_facture):
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list(
            "id_permission__code_permission",
            flat=True
        )
    )

    if "FACTURE_VALIDATE" not in permissions:
        messages.error(
            request,
            "Vous n'avez pas l'autorisation de valider une facture."
        )
        return redirect("factures")

    try:
        facture = Facture.objects.get(
            id_facture=id_facture
        )
    except Facture.DoesNotExist:
        messages.error(
            request,
            "Facture introuvable."
        )
        return redirect("factures")

    if request.method == "POST":

        if facture.statut != "EN_ATTENTE":
            messages.error(
                request,
                "Cette facture a déjÃ  été traitée."
            )
            return redirect("factures")

        facture.statut = "VALIDEE"
        facture.date_validation = timezone.now()
        facture.utilisateur_validation = str(
            request.session.get("id_utilisateur")
        )
        facture.save()

        # Validation des détails de la facture
        DetailFacture.objects.filter(
            id_facture=facture
        ).update(
            statut="VALIDEE"
        )

        messages.success(
            request,
            "Facture validée avec succès."
        )

    return redirect("factures")



def facture_detail(request, id_facture):
    if not request.session.get("id_utilisateur"):
        return redirect("connexion")

    id_utilisateur = request.session["id_utilisateur"]

    permissions = set(
        RolePermission.objects
        .filter(
            id_role__utilisateurrole__id_utilisateur=id_utilisateur,
            id_role__utilisateurrole__statut="ACTIF",
            id_permission__statut="ACTIF"
        )
        .values_list("id_permission__code_permission", flat=True)
    )

    if "FACTURE_VIEW" not in permissions:
        messages.error(request, "Vous n'avez pas l'autorisation.")
        return redirect("factures")

    try:
        facture = (
            Facture.objects
            .select_related("id_prestataire")
            .get(id_facture=id_facture)
        )
    except Facture.DoesNotExist:
        messages.error(request, "Facture introuvable.")
        return redirect("factures")

    details = (
        DetailFacture.objects
        .select_related("id_consommation", "id_acte")
        .filter(id_facture=facture)
        .order_by("id_detail_facture")
    )

    champs = get_valeurs_champs_entite("FACTURE", facture.id_facture)
    champs = [c for c in champs if c["valeur"] and str(c["valeur"]).strip()]

    return render(
        request,
        "core/facture_detail.html",
        {
            "facture": facture,
            "details": details,
            "champs_personnalises": champs,
            "page": "factures",
        }
    )



def _generer_reference_reglement(mode_reglement):
    """Génère une référence unique pour un règlement selon son mode."""
    annee = timezone.now().year

    prefixes = {
        "VIREMENT": f"VIR-{annee}-",
        "CHEQUE": f"CHQ-{annee}-",
        "ESPECES": f"ESP-{annee}-",
    }

    prefixe = prefixes.get(mode_reglement, f"REF-{annee}-")

    # Récupère toutes les références existantes pour ce préfixe
    references = (
        Reglement.objects
        .filter(reference_reglement__startswith=prefixe)
        .values_list("reference_reglement", flat=True)
    )

    valeurs = []
    for ref in references:
        try:
            valeurs.append(int(ref.rsplit("-", 1)[1]))
        except (ValueError, IndexError):
            continue

    prochain = max(valeurs, default=0) + 1
    reference = f"{prefixe}{prochain:04d}"

    # Vérifie l'unicité
    while Reglement.objects.filter(
        reference_reglement=reference
    ).exists():
        prochain += 1
        reference = f"{prefixe}{prochain:04d}"

    return reference
def _generer_numero_reglement():
    """Génère un numéro de règlement unique."""
    annee = timezone.now().year
    prefixe = f"REG-{annee}-"

    numeros = (
        Reglement.objects
        .filter(numero_reglement__startswith=prefixe)
        .values_list("numero_reglement", flat=True)
    )

